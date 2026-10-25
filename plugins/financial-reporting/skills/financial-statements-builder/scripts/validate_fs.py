#!/usr/bin/env python3
"""
validate_fs.py — hard-gate validation of a financial statements workbook.

Reads the workbook with cached values (run a LibreOffice recalculation first, e.g.
`soffice --headless --convert-to xlsx --outdir recalc/ book.xlsx`) and runs:

  C01/C06  TB sums to zero; mapping total = TB total          (needs TB_Raw, Mapping tabs)
  C04/C05  every ledger mapped exactly once                    (Mapping tab)
  C07      balance sheet balances                              (named cells or Checks tab)
  C08      P&L -> retained earnings link
  C10      cash flow closing cash = BS cash
  C13      every note total = face line                        (Checks tab rows "C13*")
  C19      no Excel error values anywhere
  ALL      any row on the Checks tab whose Status is FAIL

Workbook conventions expected (the skill builds them this way):
  - Tab "Mapping" with headers: ledger_code, ledger_name, closing_balance, fsli, note_no
  - Tab "TB_Raw"  with header closing_balance
  - Tab "Checks"  with headers: Check ID | Check | Severity | Expected | Actual | Difference | Status
  - Optional defined names: TOTAL_ASSETS_CY, TOTAL_EQ_LIAB_CY, CF_CLOSING_CASH_CY, BS_CASH_CY,
    RE_CLOSING_CY, RE_OPENING_CY, PAT_CY, DIVIDENDS_CY, RE_OTHER_CY (and _PY equivalents)

Usage: python validate_fs.py workbook.xlsx [--tol 0.5]
Exit code 0 = all CRITICAL passed; 1 = at least one CRITICAL failure.
"""
import argparse
import sys
from collections import Counter

from openpyxl import load_workbook

ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?", "#NUM!", "#NULL!")


def header_index(ws):
    hdr = [str(c.value).strip().lower() if c.value is not None else "" for c in ws[1]]
    return {h: i for i, h in enumerate(hdr) if h}


def col_values(ws, name):
    idx = header_index(ws)
    if name not in idx:
        return None
    return [row[idx[name]] for row in ws.iter_rows(min_row=2, values_only=True)]


def named(wb, name):
    dn = wb.defined_names.get(name) if hasattr(wb.defined_names, "get") else None
    if dn is None:
        return None
    for sheet, ref in dn.destinations:
        v = wb[sheet][ref.replace("$", "")].value
        return float(v) if isinstance(v, (int, float)) else None
    return None


def num(x):
    try:
        return float(x)
    except (TypeError, ValueError):
        return 0.0


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook")
    ap.add_argument("--tol", type=float, default=0.5)
    a = ap.parse_args()

    wb = load_workbook(a.workbook, data_only=True)
    results = []  # (id, check, severity, diff, passed)

    def add(cid, text, sev, diff, ok=None):
        passed = ok if ok is not None else (diff is not None and abs(diff) <= a.tol)
        results.append((cid, text, sev, diff, passed))

    # C19 — error values
    errs = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.strip() in ERRORS:
                    errs.append(f"{ws.title}!{c.coordinate}")
    add("C19", f"No formula errors ({len(errs)} found)", "CRITICAL", None, ok=not errs)

    # C01 / C04 / C05 / C06
    tb_total = None
    if "TB_Raw" in wb.sheetnames:
        v = col_values(wb["TB_Raw"], "closing_balance")
        if v is not None:
            tb_total = sum(num(x) for x in v)
            add("C01", "Trial balance sums to zero", "CRITICAL", tb_total)
    if "Mapping" in wb.sheetnames:
        ws = wb["Mapping"]
        codes = col_values(ws, "ledger_code") or col_values(ws, "ledger_name") or []
        fsli = col_values(ws, "fsli") or []
        bal = col_values(ws, "closing_balance") or []
        unmapped = sum(1 for f in fsli if f in (None, "") or str(f).startswith("SUSPENSE"))
        dup = [k for k, n in Counter(c for c in codes if c not in (None, "")).items() if n > 1]
        add("C04", f"All ledgers mapped ({unmapped} unmapped/suspense)", "CRITICAL", None, ok=unmapped == 0)
        add("C05", f"No double-mapped ledgers ({len(dup)} duplicates)", "CRITICAL", None, ok=not dup)
        if tb_total is not None:
            add("C06", "Mapping total = TB total", "CRITICAL", sum(num(x) for x in bal) - tb_total)

    # C07 / C08 / C10 via defined names
    for p in ("CY", "PY"):
        ta, tel = named(wb, f"TOTAL_ASSETS_{p}"), named(wb, f"TOTAL_EQ_LIAB_{p}")
        if ta is not None and tel is not None:
            add(f"C07-{p}", f"Balance sheet balances ({p})", "CRITICAL", ta - tel)
        cf, bs = named(wb, f"CF_CLOSING_CASH_{p}"), named(wb, f"BS_CASH_{p}")
        if cf is not None and bs is not None:
            add(f"C10-{p}", f"Cash flow closing cash = BS cash ({p})", "CRITICAL", cf - bs)
        rc, ro, pat = named(wb, f"RE_CLOSING_{p}"), named(wb, f"RE_OPENING_{p}"), named(wb, f"PAT_{p}")
        if None not in (rc, ro, pat):
            div = named(wb, f"DIVIDENDS_{p}") or 0
            oth = named(wb, f"RE_OTHER_{p}") or 0
            add(f"C08-{p}", f"P&L to retained earnings link ({p})", "CRITICAL", rc - (ro + pat - div + oth))

    # Checks tab — trust nothing, re-read status and difference
    if "Checks" in wb.sheetnames:
        ws = wb["Checks"]
        idx = header_index(ws)
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[0] in (None, ""):
                continue
            cid = str(row[idx.get("check id", 0)])
            text = str(row[idx.get("check", 1)])
            sev = str(row[idx.get("severity", 2)] or "CRITICAL").upper()
            diff = row[idx["difference"]] if "difference" in idx else None
            status = str(row[idx["status"]]).upper() if "status" in idx else ""
            ok = status == "PASS" and (diff is None or abs(num(diff)) <= a.tol)
            add(cid, text, sev, num(diff) if diff is not None else None, ok=ok)

    # Report
    w = max(len(r[1]) for r in results) if results else 20
    print(f"{'ID':<8} {'Check':<{w}} {'Severity':<9} {'Difference':>14}  Status")
    for cid, text, sev, diff, ok in results:
        d = f"{diff:,.2f}" if isinstance(diff, (int, float)) else "-"
        print(f"{cid:<8} {text:<{w}} {sev:<9} {d:>14}  {'PASS' if ok else 'FAIL'}")
    if errs:
        print("Error cells (first 20):", ", ".join(errs[:20]))
    crit_fail = [r for r in results if r[2] == "CRITICAL" and not r[4]]
    warn_fail = [r for r in results if r[2] != "CRITICAL" and not r[4]]
    print(f"\nSummary: {len(results) - len(crit_fail) - len(warn_fail)}/{len(results)} passed | "
          f"CRITICAL failures: {len(crit_fail)} | WARNING failures: {len(warn_fail)}")
    if not results:
        print("No checks could be run — workbook does not follow the expected conventions.")
        sys.exit(1)
    sys.exit(1 if crit_fail else 0)


if __name__ == "__main__":
    main()
